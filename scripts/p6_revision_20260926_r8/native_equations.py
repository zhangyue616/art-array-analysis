"""Optional helper for inserting the published MathML definitions into Word.

This module is separate from the offline reproduction route. It requires
``python-docx``, ``lxml``, and a local Microsoft Office ``MML2OMML.XSL`` file.
Pass the transform explicitly or set ``ART_ARRAY_MML2OMML_XSL``.
"""

from __future__ import annotations

import copy
import json
import os
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SPECS = Path(__file__).with_name("equations.json")


def resolve_mml2omml_xsl(requested: str | Path | None = None) -> Path:
    """Resolve Microsoft's local MathML-to-OMML transform on Windows."""
    candidates: list[Path] = []
    if requested is not None:
        candidates.append(Path(requested))
    configured = os.environ.get("ART_ARRAY_MML2OMML_XSL")
    if configured:
        candidates.append(Path(configured))
    for variable in ("ProgramFiles", "ProgramFiles(x86)"):
        base = os.environ.get(variable)
        if base:
            candidates.append(
                Path(base) / "Microsoft Office" / "root" / "Office16" / "MML2OMML.XSL"
            )
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    raise FileNotFoundError(
        "Microsoft Office MML2OMML.XSL was not found. Pass xsl_path or set "
        "ART_ARRAY_MML2OMML_XSL."
    )


def append_equation(document, identifier: str, *, xsl_path: str | Path | None = None):
    """Append one centered editable Word equation and return its saved spec."""
    spec = json.loads(SPECS.read_text(encoding="utf-8"))[identifier]
    transform = etree.XSLT(etree.parse(str(resolve_mml2omml_xsl(xsl_path))))
    mathml = etree.fromstring(spec["mathml"].encode("utf-8"))
    transformed = transform(mathml).getroot()
    native = (
        transformed
        if transformed.tag == qn("m:oMath")
        else transformed.find(".//" + qn("m:oMath"))
    )
    if native is None:
        raise ValueError("MathML conversion did not produce native Word math: " + identifier)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    section = document.sections[-1]
    usable = section.page_width - section.left_margin - section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        int(usable / 2), WD_TAB_ALIGNMENT.CENTER
    )
    paragraph.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    paragraph._p.append(copy.deepcopy(native))
    paragraph.add_run(".\t(" + spec["number"] + ")")
    return spec
